"""analysis_xlsx — a fully formula-driven analysis workbook (show all the equations).

Sheet 1 "Statement Data": the raw audited lines (prior + current FY, in $m) as inputs.
Sheet 2 "Calculations": every ratio, DuPont, ROIC vs WACC, Piotroski signal, Altman Z, Beneish M
and Sloan accrual computed with LIVE Excel formulas that reference Sheet 1 — change a data cell and
everything recalculates. Divisions are wrapped in IFERROR so the workbook never shows error tokens.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUTDIR = Path(__file__).resolve().parent / "output"
M = 1_000_000
BOLD = Font(bold=True)
HEAD = Font(bold=True, color="FFFFFF")
NAVY = PatternFill("solid", fgColor="0F172A")
SECT = PatternFill("solid", fgColor="E2E8F0")
TITLE = Font(bold=True, size=14)
PCT = '0.0%'
NUM2 = '0.00'

# Statement-data rows on the Data sheet (label, fund-concept, is_monetary)
DATA_ROWS = [
    ("Revenue", "revenue", True), ("Cost of goods sold", "cogs", True),
    ("Gross profit", "gross_profit", True), ("Operating income (EBIT)", "operating_income", True),
    ("Depreciation & amortisation", "da", True), ("SG&A", "sga", True),
    ("Net income", "net_income", True), ("Cash from operations", "cfo", True),
    ("Capex", "capex", True), ("Total assets", "total_assets", True),
    ("Current assets", "current_assets", True), ("Current liabilities", "current_liabilities", True),
    ("Total liabilities", "total_liabilities", True), ("Shareholders' equity", "equity", True),
    ("Retained earnings", "retained_earnings", True), ("Receivables", "receivables", True),
    ("PP&E (net)", "ppe_net", True), ("Long-term debt", "long_term_debt", True),
    ("Short-term debt", "short_term_debt", True),
    ("Cash & equivalents", "cash", True), ("Diluted shares (m)", "shares", True),
]


def build(fund: dict, rep: dict) -> Path:
    years = fund["years"]
    if len(years) < 2:
        return _single_year_fallback(fund, rep)
    yp, yc = years[-2], years[-1]
    wb = Workbook()

    # ---------------- Statement Data ----------------
    D = wb.active; D.title = "Statement Data"
    D.column_dimensions["A"].width = 30
    D.column_dimensions["B"].width = 15; D.column_dimensions["C"].width = 15
    D["A1"] = f"{fund['name']} ({fund['ticker']}) — Statement data ($m)"; D["A1"].font = TITLE
    for col, yr in (("B", yp), ("C", yc)):
        D[f"{col}2"] = f"FY{yr}"; D[f"{col}2"].font = HEAD; D[f"{col}2"].fill = NAVY
    rowof = {}
    for i, (label, concept, _) in enumerate(DATA_ROWS):
        r = 3 + i
        rowof[concept] = r
        D[f"A{r}"] = label
        for col, yr in (("B", yp), ("C", yc)):
            val = fund.get(concept, {}).get(yr)
            D[f"{col}{r}"] = (val / M) if val is not None else None
            D[f"{col}{r}"].number_format = '#,##0;(#,##0)'
    # input block
    ir = 3 + len(DATA_ROWS) + 1
    D[f"A{ir}"] = "Market cap ($m)"; D[f"B{ir}"] = (rep.get("market_cap") or 0) / M
    D[f"A{ir+1}"] = "WACC"; D[f"B{ir+1}"] = rep.get("wacc", 0.10); D[f"B{ir+1}"].number_format = PCT
    D[f"A{ir+2}"] = "Tax rate"; D[f"B{ir+2}"] = 0.21; D[f"B{ir+2}"].number_format = PCT
    mc = f"'Statement Data'!B{ir}"
    wacc_c = f"'Statement Data'!B{ir+1}"
    tax_c = f"'Statement Data'!B{ir+2}"

    def c(concept):  # current-year data cell
        return f"'Statement Data'!C{rowof[concept]}"

    def p(concept):  # prior-year data cell
        return f"'Statement Data'!B{rowof[concept]}"

    def mref(cell):
        return f"'Statement Data'!{cell}"

    # ---------------- Calculations ----------------
    C = wb.create_sheet("Calculations")
    C.column_dimensions["A"].width = 34; C.column_dimensions["B"].width = 14
    C.column_dimensions["C"].width = 40
    C["A1"] = "Calculations (live formulas → Statement Data)"; C["A1"].font = TITLE
    row = 3

    def section(name):
        nonlocal row
        C[f"A{row}"] = name; C[f"A{row}"].font = BOLD; C[f"A{row}"].fill = SECT
        C[f"B{row}"].fill = SECT; C[f"C{row}"].fill = SECT
        row += 1

    def line(label, formula, fmt=NUM2, note=""):
        nonlocal row
        cell = f"B{row}"
        C[f"A{row}"] = label
        C[cell] = formula
        C[cell].number_format = fmt
        if note:
            C[f"C{row}"] = note; C[f"C{row}"].font = Font(italic=True, color="64748B")
        row += 1
        return cell

    section("PROFITABILITY")
    line("Gross margin", f"=IFERROR({c('gross_profit')}/{c('revenue')},\"n/a\")", PCT, "Gross profit / Revenue")
    line("Operating margin", f"=IFERROR({c('operating_income')}/{c('revenue')},\"n/a\")", PCT, "EBIT / Revenue")
    line("Net margin", f"=IFERROR({c('net_income')}/{c('revenue')},\"n/a\")", PCT, "Net income / Revenue")
    line("ROA", f"=IFERROR({c('net_income')}/{c('total_assets')},\"n/a\")", PCT, "Net income / Total assets")
    line("ROE", f"=IFERROR({c('net_income')}/{c('equity')},\"n/a\")", PCT, "Net income / Equity")

    section("DUPONT  (ROE = margin x turnover x leverage)")
    nm = line("Net margin", f"=IFERROR({c('net_income')}/{c('revenue')},\"n/a\")", PCT)
    at = line("Asset turnover", f"=IFERROR({c('revenue')}/{c('total_assets')},\"n/a\")", NUM2, "Revenue / Total assets")
    em = line("Equity multiplier", f"=IFERROR({c('total_assets')}/{c('equity')},\"n/a\")", NUM2, "Total assets / Equity")
    line("ROE (DuPont)", f"=IFERROR({nm}*{at}*{em},\"n/a\")", PCT, "= the three above multiplied")

    section("VALUE CREATION")
    roic = line("ROIC", f"=IFERROR({c('operating_income')}*(1-{tax_c})/({c('long_term_debt')}+{c('short_term_debt')}+{c('equity')}-{c('cash')}),\"n/a\")",
                PCT, "NOPAT / (total debt + equity - cash)")
    line("WACC", f"={wacc_c}", PCT, "cost of capital (input)")
    line("ROIC - WACC spread", f"=IFERROR({roic}-{wacc_c},\"n/a\")", PCT, ">0 = creates value")

    section("BALANCE-SHEET HEALTH")
    line("Current ratio", f"=IFERROR({c('current_assets')}/{c('current_liabilities')},\"n/a\")", NUM2)
    line("LT debt / Equity", f"=IFERROR({c('long_term_debt')}/{c('equity')},\"n/a\")", NUM2)
    line("Net debt / EBITDA", f"=IFERROR(({c('long_term_debt')}+{c('short_term_debt')}-{c('cash')})/({c('operating_income')}+{c('da')}),\"n/a\")", NUM2)

    section("PIOTROSKI F-SCORE  (1 point each)")
    sig_rows = []
    def signal(label, formula, note):
        nonlocal row
        sig_rows.append(f"B{row}")
        line(label, formula, "0", note)
    signal("Positive ROA", f"=IFERROR(IF({c('net_income')}/{c('total_assets')}>0,1,0),0)", "profitable on assets")
    signal("Positive CFO", f"=IF({c('cfo')}>0,1,0)", "operating cash positive")
    signal("Rising ROA", f"=IFERROR(IF({c('net_income')}/{c('total_assets')}>{p('net_income')}/{p('total_assets')},1,0),0)", "")
    signal("CFO > Net income", f"=IFERROR(IF({c('cfo')}/{c('total_assets')}>{c('net_income')}/{c('total_assets')},1,0),0)", "accrual quality")
    signal("Lower leverage", f"=IFERROR(IF({c('long_term_debt')}/{c('total_assets')}<{p('long_term_debt')}/{p('total_assets')},1,0),0)", "")
    signal("Higher current ratio", f"=IFERROR(IF({c('current_assets')}/{c('current_liabilities')}>{p('current_assets')}/{p('current_liabilities')},1,0),0)", "")
    signal("No dilution", f"=IFERROR(IF({c('shares')}<={p('shares')}*1.001,1,0),0)", "shares not up")
    signal("Higher gross margin", f"=IFERROR(IF({c('gross_profit')}/{c('revenue')}>{p('gross_profit')}/{p('revenue')},1,0),0)", "")
    signal("Higher asset turnover", f"=IFERROR(IF({c('revenue')}/{c('total_assets')}>{p('revenue')}/{p('total_assets')},1,0),0)", "")
    line("F-Score (/9)", f"=SUM({sig_rows[0]}:{sig_rows[-1]})", "0", "7-9 strong, 0-3 weak")
    C[f"A{row-1}"].font = BOLD

    section("ALTMAN Z-SCORE")
    line("Altman Z", (f"=IFERROR(1.2*({c('current_assets')}-{c('current_liabilities')})/{c('total_assets')}"
                      f"+1.4*{c('retained_earnings')}/{c('total_assets')}+3.3*{c('operating_income')}/{c('total_assets')}"
                      f"+0.6*{mc}/{c('total_liabilities')}+{c('revenue')}/{c('total_assets')},\"n/a\")"),
         NUM2, ">2.99 safe, <1.81 distress")

    section("BENEISH M-SCORE  (>-1.78 = possible manipulation)")
    dsri = line("DSRI (receivables)", f"=IFERROR(({c('receivables')}/{c('revenue')})/({p('receivables')}/{p('revenue')}),\"n/a\")")
    gmi = line("GMI (gross margin)", f"=IFERROR(({p('gross_profit')}/{p('revenue')})/({c('gross_profit')}/{c('revenue')}),\"n/a\")")
    aqi = line("AQI (asset quality)", f"=IFERROR((1-({c('current_assets')}+{c('ppe_net')})/{c('total_assets')})/(1-({p('current_assets')}+{p('ppe_net')})/{p('total_assets')}),\"n/a\")")
    sgi = line("SGI (sales growth)", f"=IFERROR({c('revenue')}/{p('revenue')},\"n/a\")")
    depi = line("DEPI (depreciation)", f"=IFERROR(({p('da')}/({p('da')}+{p('ppe_net')}))/({c('da')}/({c('da')}+{c('ppe_net')})),\"n/a\")")
    sgai = line("SGAI (SG&A)", f"=IFERROR(({c('sga')}/{c('revenue')})/({p('sga')}/{p('revenue')}),\"n/a\")")
    lvgi = line("LVGI (leverage)", f"=IFERROR(({c('total_liabilities')}/{c('total_assets')})/({p('total_liabilities')}/{p('total_assets')}),\"n/a\")")
    tata = line("TATA (accruals)", f"=IFERROR(({c('net_income')}-{c('cfo')})/{c('total_assets')},\"n/a\")")
    line("M-Score", f"=IFERROR(-4.84+0.92*{dsri}+0.528*{gmi}+0.404*{aqi}+0.892*{sgi}+0.115*{depi}-0.172*{sgai}+4.679*{tata}-0.327*{lvgi},\"n/a\")",
         NUM2, "Beneish 1999")
    C[f"A{row-1}"].font = BOLD

    section("EARNINGS QUALITY")
    line("Sloan accrual ratio", f"=IFERROR(({c('net_income')}-{c('cfo')})/{c('total_assets')},\"n/a\")", NUM2, ">0.1 low quality")

    section("VERDICT")
    rec = rep.get("recommendation", {})
    line("Recommendation", f"\"{rec.get('rating','n/a')}  (score {rec.get('score','')})\"", "General")
    rd = rep.get("reverse_dcf", {}).get("implied_growth")
    line("Reverse-DCF implied growth", (rd if rd is not None else "n/a"), PCT, "growth the price implies")

    OUTDIR.mkdir(exist_ok=True)
    out = OUTDIR / f"{fund['ticker']}_analysis_workbook.xlsx"
    wb.save(out)
    return out


def _single_year_fallback(fund: dict, rep: dict) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "Calculations"
    ws["A1"] = f"{fund['ticker']}: only one fiscal year available — trend-based scores need two."
    OUTDIR.mkdir(exist_ok=True)
    out = OUTDIR / f"{fund['ticker']}_analysis_workbook.xlsx"
    wb.save(out)
    return out
