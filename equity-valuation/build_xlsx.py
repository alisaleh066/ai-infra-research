"""build_xlsx — build a linked DCF workbook for ANY ticker.

Conventions: blue = inputs, black = formulas, dedicated Assumptions tab, units in headers,
parentheses for negatives, zero formula errors. Sheets: Assumptions, Model (income statement +
FCFF projection), WACC, DCF, Sensitivity, Comps. All monetary values in USD millions.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTDIR = Path(__file__).resolve().parent / "output"

BLUE = Font(color="1D4ED8")
BLACK = Font(color="0F172A")
BOLD = Font(color="0F172A", bold=True)
HEADWHITE = Font(color="FFFFFF", bold=True)
NAVY = PatternFill("solid", fgColor="0F172A")
SLATE = PatternFill("solid", fgColor="E2E8F0")
TITLE = Font(color="0F172A", bold=True, size=14)
ITAL = Font(italic=True, color="64748B")
MONEY = '#,##0;(#,##0)'
PCT = '0.0%'
PS = '"$"#,##0.00'
NUM2 = '0.00'
M = 1_000_000


def _inp(ws, c, v, fmt=None):
    ws[c].value = v; ws[c].font = BLUE
    if fmt: ws[c].number_format = fmt


def _fml(ws, c, v, fmt=None):
    ws[c].value = v; ws[c].font = BLACK
    if fmt: ws[c].number_format = fmt


def _lbl(ws, c, t, bold=False):
    ws[c].value = t; ws[c].font = BOLD if bold else BLACK


def _hdr(ws, c, t):
    ws[c].value = t; ws[c].font = HEADWHITE; ws[c].fill = NAVY
    ws[c].alignment = Alignment(horizontal="center")


def _mm(series: dict, year: int):
    """Value in $m for a fiscal year, or None."""
    v = series.get(year)
    return None if v is None else v / M


def build(fund, a, dcf, comps, waccs, growths, grid, summary, extras) -> Path:
    years = fund["years"]
    H = len(years)
    hist = [get_column_letter(2 + i) for i in range(H)]          # historical cols
    proj = [get_column_letter(2 + H + i) for i in range(5)]      # projection cols
    last_h = hist[-1]
    wb = Workbook()

    # --------------------------------------------------------- Assumptions
    A = wb.active; A.title = "Assumptions"
    A.column_dimensions["A"].width = 36
    for col in "BCDEFGHIJKL":
        A.column_dimensions[col].width = 12
    A["A1"].value = f"{fund['name']} ({fund['ticker']}) — Valuation Assumptions"; A["A1"].font = TITLE
    A["A2"].value = "Blue = input (editable). Fundamentals from SEC EDGAR; market/assumptions flagged."
    A["A2"].font = ITAL
    rows = [
        ("Tax rate [ASSUMPTION]", a.tax_rate, PCT),
        ("Risk-free rate [ASSUMPTION]", a.risk_free, PCT),
        ("Equity risk premium [ASSUMPTION]", a.erp, PCT),
        ("Beta", round(a.beta, 3), NUM2),
        ("Pre-tax cost of debt [ASSUMPTION]", a.pretax_kd, PCT),
        ("Share price (USD)", summary["price"] or 0, PS),
        ("Shares outstanding (m)", a.shares, MONEY),
        ("Terminal growth [ASSUMPTION]", a.terminal_growth, PCT),
        ("D&A (% revenue)", a.da_pct_rev, PCT),
        ("Capex (% revenue)", a.capex_pct_rev, PCT),
        ("Incr. NWC (% of rev change) [ASSUMPTION]", a.nwc_pct_rev_change, PCT),
        ("Total debt ($m) — EDGAR", a.debt_value, MONEY),
        ("Cash ($m) — EDGAR", a.cash, MONEY),
    ]
    r = {}
    for i, (lbl, val, fmt) in enumerate(rows):
        rr = 4 + i
        _lbl(A, f"A{rr}", lbl); _inp(A, f"B{rr}", val, fmt); r[lbl.split(" [")[0]] = rr
    g_row0 = 4 + len(rows) + 1
    _lbl(A, f"A{g_row0}", "Revenue growth Yr1–Yr5 [ASSUMPTION]", bold=True)
    for i, g in enumerate(a.rev_growth):
        _lbl(A, f"A{g_row0+1+i}", f"  growth Yr{i+1}"); _inp(A, f"B{g_row0+1+i}", g, PCT)
    m_row0 = g_row0 + 6 + 1
    _lbl(A, f"A{m_row0}", "EBIT margin Yr1–Yr5 [ASSUMPTION]", bold=True)
    for i, m in enumerate(a.ebit_margin):
        _lbl(A, f"A{m_row0+1+i}", f"  margin Yr{i+1}"); _inp(A, f"B{m_row0+1+i}", m, PCT)
    growth_cells = [f"B{g_row0+1+i}" for i in range(5)]
    margin_cells = [f"B{m_row0+1+i}" for i in range(5)]
    A_tax = f"Assumptions!B{r['Tax rate']}"
    A_da = f"Assumptions!B{r['D&A (% revenue)']}"
    A_capex = f"Assumptions!B{r['Capex (% revenue)']}"
    A_nwc = f"Assumptions!B{r['Incr. NWC (% of rev change)']}"
    A_tg = f"Assumptions!B{r['Terminal growth']}"
    A_debt = f"Assumptions!B{r['Total debt ($m) — EDGAR']}"
    A_cash = f"Assumptions!B{r['Cash ($m) — EDGAR']}"
    A_shares = f"Assumptions!B{r['Shares outstanding (m)']}"

    # --------------------------------------------------------- Model
    Mo = wb.create_sheet("Model")
    Mo.column_dimensions["A"].width = 30
    for i in range(H + 5):
        Mo.column_dimensions[get_column_letter(2 + i)].width = 12
    Mo["A1"].value = f"{fund['ticker']} — Income statement & FCFF ($m)"; Mo["A1"].font = TITLE
    for i, y in enumerate(years):
        _hdr(Mo, f"{hist[i]}3", f"FY{y}")
    for i in range(5):
        _hdr(Mo, f"{proj[i]}3", f"Yr{i+1}")

    _lbl(Mo, "A4", "Revenue ($m)", bold=True)
    for i, y in enumerate(years):
        _inp(Mo, f"{hist[i]}4", _mm(fund["revenue"], y), MONEY)
    prev = last_h
    for i in range(5):
        _fml(Mo, f"{proj[i]}4", f"={prev}4*(1+{growth_cells[i].replace('B', 'Assumptions!B')})", MONEY)
        prev = proj[i]
    _lbl(Mo, "A5", "Operating income / EBIT ($m)", bold=True)
    for i, y in enumerate(years):
        _inp(Mo, f"{hist[i]}5", _mm(fund["operating_income"], y), MONEY)
    for i in range(5):
        _fml(Mo, f"{proj[i]}5", f"={proj[i]}4*{margin_cells[i].replace('B', 'Assumptions!B')}", MONEY)
    _lbl(Mo, "A6", "EBIT margin (%)")
    for c in hist + proj:
        _fml(Mo, f"{c}6", f"={c}5/{c}4", PCT)
    _lbl(Mo, "A7", "Net income ($m)")
    for i, y in enumerate(years):
        _inp(Mo, f"{hist[i]}7", _mm(fund["net_income"], y), MONEY)
    _lbl(Mo, "A8", "EPS diluted ($)")
    for i, y in enumerate(years):
        _inp(Mo, f"{hist[i]}8", fund["eps_diluted"].get(y), NUM2)

    _lbl(Mo, "A10", "Free cash flow to firm (projection)", bold=True)
    _lbl(Mo, "A11", "NOPAT ($m)")
    for i in range(5):
        _fml(Mo, f"{proj[i]}11", f"={proj[i]}5*(1-{A_tax})", MONEY)
    _lbl(Mo, "A12", "D&A ($m)")
    for i in range(5):
        _fml(Mo, f"{proj[i]}12", f"={proj[i]}4*{A_da}", MONEY)
    _lbl(Mo, "A13", "Capex ($m)")
    for i in range(5):
        _fml(Mo, f"{proj[i]}13", f"={proj[i]}4*{A_capex}", MONEY)
    _lbl(Mo, "A14", "Incr. NWC ($m)")
    prev = last_h
    for i in range(5):
        _fml(Mo, f"{proj[i]}14", f"=({proj[i]}4-{prev}4)*{A_nwc}", MONEY); prev = proj[i]
    _lbl(Mo, "A15", "FCFF ($m)", bold=True)
    for i in range(5):
        _fml(Mo, f"{proj[i]}15", f"={proj[i]}11+{proj[i]}12-{proj[i]}13-{proj[i]}14", MONEY)

    # --------------------------------------------------------- WACC
    W = wb.create_sheet("WACC")
    W.column_dimensions["A"].width = 26; W.column_dimensions["B"].width = 14
    W["A1"].value = "WACC build"; W["A1"].font = TITLE
    A_rf = f"Assumptions!B{r['Risk-free rate']}"; A_beta = f"Assumptions!B{r['Beta']}"
    A_erp = f"Assumptions!B{r['Equity risk premium']}"; A_kd = f"Assumptions!B{r['Pre-tax cost of debt']}"
    A_price = f"Assumptions!B{r['Share price (USD)']}"
    pairs = [
        ("Cost of equity (CAPM)", f"={A_rf}+{A_beta}*{A_erp}", PCT, "B4"),
        ("After-tax cost of debt", f"={A_kd}*(1-{A_tax})", PCT, "B5"),
        ("Equity value ($m)", f"={A_price}*{A_shares}", MONEY, "B6"),
        ("Debt value ($m)", f"={A_debt}", MONEY, "B7"),
        ("Equity weight", "=B6/(B6+B7)", PCT, "B8"),
        ("Debt weight", "=B7/(B6+B7)", PCT, "B9"),
        ("WACC", "=B8*B4+B9*B5", PCT, "B10"),
    ]
    for i, (lbl, f_, fmt, cell) in enumerate(pairs):
        _lbl(W, f"A{4+i}", lbl); _fml(W, cell, f_, fmt)
    W["A10"].font = BOLD; W["B10"].font = BOLD

    # --------------------------------------------------------- DCF
    D = wb.create_sheet("DCF")
    D.column_dimensions["A"].width = 30
    for c in "BCDEF":
        D.column_dimensions[c].width = 13
    D["A1"].value = "Discounted cash flow (FCFF)"; D["A1"].font = TITLE
    _lbl(D, "A3", "Year")
    for i, c in enumerate("BCDEF"):
        D[f"{c}3"].value = i + 1; D[f"{c}3"].font = BLACK
    _lbl(D, "A4", "FCFF ($m)")
    for i, c in enumerate("BCDEF"):
        _fml(D, f"{c}4", f"=Model!{proj[i]}15", MONEY)
    _lbl(D, "A5", "Discount factor")
    for c in "BCDEF":
        _fml(D, f"{c}5", f"=1/(1+WACC!$B$10)^{c}3", NUM2)
    _lbl(D, "A6", "PV of FCFF ($m)")
    for c in "BCDEF":
        _fml(D, f"{c}6", f"={c}4*{c}5", MONEY)
    _lbl(D, "A8", "Sum PV of FCFF ($m)"); _fml(D, "B8", "=SUM(B6:F6)", MONEY)
    _lbl(D, "A9", "Terminal value ($m)"); _fml(D, "B9", f"=F4*(1+{A_tg})/(WACC!B10-{A_tg})", MONEY)
    _lbl(D, "A10", "PV of terminal value ($m)"); _fml(D, "B10", "=B9*F5", MONEY)
    _lbl(D, "A11", "Enterprise value ($m)", bold=True); _fml(D, "B11", "=B8+B10", MONEY)
    _lbl(D, "A12", "Less: net debt ($m)"); _fml(D, "B12", f"={A_debt}-{A_cash}", MONEY)
    _lbl(D, "A13", "Equity value ($m)", bold=True); _fml(D, "B13", "=B11-B12", MONEY)
    _lbl(D, "A14", "Shares (m)"); _fml(D, "B14", f"={A_shares}", MONEY)
    _lbl(D, "A15", "Value per share ($)", bold=True); _fml(D, "B15", "=B13/B14", PS)
    D["B15"].fill = SLATE
    _lbl(D, "A16", "TV as % of EV"); _fml(D, "B16", "=B10/B11", PCT)

    # --------------------------------------------------------- Sensitivity
    S = wb.create_sheet("Sensitivity")
    S.column_dimensions["A"].width = 12
    for c in "BCDEFG":
        S.column_dimensions[c].width = 12
    S["A1"].value = "Value per share — WACC (rows) × terminal growth (cols)"; S["A1"].font = TITLE
    _lbl(S, "B3", "WACC\\g", bold=True)
    for j, g in enumerate(growths):
        _inp(S, f"{chr(67+j)}3", g, PCT)
    deltas = [-0.015, -0.0075, 0.0, 0.0075, 0.015]
    for i, dval in enumerate(deltas):
        rr = 4 + i
        _fml(S, f"B{rr}", "=WACC!$B$10" + (f"+{dval}" if dval > 0 else (f"{dval}" if dval < 0 else "")), PCT)
        for j in range(5):
            col = chr(67 + j)
            formula = (f"=(SUMPRODUCT(DCF!$B$4:$F$4,1/(1+$B{rr})^DCF!$B$3:$F$3)"
                       f"+(DCF!$F$4*(1+{col}$3)/($B{rr}-{col}$3))/(1+$B{rr})^5"
                       f"-({A_debt}-{A_cash}))/{A_shares}")
            _fml(S, f"{col}{rr}", formula, PS)
    S["E6"].fill = SLATE

    # --------------------------------------------------------- Comps
    C = wb.create_sheet("Comps")
    C.column_dimensions["A"].width = 36; C.column_dimensions["B"].width = 14
    C["A1"].value = "Comparable-company valuation"; C["A1"].font = TITLE
    C["A2"].value = "Peer multiples are ASSUMPTION inputs — replace with live peers."; C["A2"].font = ITAL
    # EBITDA = latest EBIT + latest D&A ($m); D&A not shown historically on Model, so input here.
    _lbl(C, "A4", "EBITDA ($m)"); _inp(C, "B4", round(extras["ebitda"], 1), MONEY)
    _lbl(C, "A5", "EV/EBITDA multiple (x) [ASSUMPTION]"); _inp(C, "B5", 12.0, NUM2)
    _lbl(C, "A6", "Implied EV ($m)"); _fml(C, "B6", "=B4*B5", MONEY)
    _lbl(C, "A7", "Less: net debt ($m)"); _fml(C, "B7", f"={A_debt}-{A_cash}", MONEY)
    _lbl(C, "A8", "Equity value ($m)"); _fml(C, "B8", "=B6-B7", MONEY)
    _lbl(C, "A9", "Value/share — EV/EBITDA ($)"); _fml(C, "B9", f"=B8/{A_shares}", PS)
    _lbl(C, "A11", "EPS ($)"); _inp(C, "B11", round(extras["eps"] or 0.0, 2), NUM2)
    _lbl(C, "A12", "P/E multiple (x) [ASSUMPTION]"); _inp(C, "B12", 18.0, NUM2)
    _lbl(C, "A13", "Value/share — P/E ($)"); _fml(C, "B13", "=B11*B12", PS)
    _lbl(C, "A14", "Comps midpoint ($)", bold=True); _fml(C, "B14", "=AVERAGE(B9,B13)", PS)
    C["B14"].fill = SLATE

    OUTDIR.mkdir(exist_ok=True)
    out = OUTDIR / f"{fund['ticker']}_valuation.xlsx"
    wb.save(out)
    return out
